"""Tests for the parts of the ETL sources that don't need a download: shop and road rules, and sheet parsers."""
import pandas as pd
import pytest
from openpyxl import Workbook

import geopandas as gpd
from shapely.geometry import Point

from etl.core import EtlError
from etl.sources import abs as abs_sources
from etl.sources import cityofsydney, osm


# --- OpenStreetMap: shops and roads --------------------------------------------------------------
@pytest.mark.parametrize("shop, amenity, text, expected", [
    ("supermarket", None, "Harris Farm", "supermarket"),
    ("convenience", None, "IGA Express Balmain", "supermarket"),  # a chain, so a supermarket in the AUO list
    ("grocery", None, "Woolworths Metro", "supermarket"),
    ("convenience", None, "Local Milk Bar", "convenience"),
    ("newsagent", None, "Central Newsagency", "convenience"),
    (None, "fuel", "Coles Express", "convenience"),  # a petrol station, not a Coles supermarket
    ("greengrocer", None, "Fruit Barn", None),
    ("grocery", None, "Continental Deli", None),
    ("kiosk", None, "", None)])
def test_shops_follow_the_auo_daily_living_definitions(shop, amenity, text, expected):
    assert osm.classify_shop(shop, amenity, text) == expected


def test_a_word_inside_another_word_is_not_a_chain_name():
    # "iga" sits inside "Bigalow", but a shop with that name is an ordinary convenience store
    assert osm.classify_shop("convenience", None, "Bigalow Mini Mart") == "convenience"


@pytest.mark.parametrize("tags, expected", [
    ({}, True), ({"foot": "no"}, False), ({"motorroad": "yes"}, False),
    ({"access": "private"}, False), ({"access": "private", "foot": "yes"}, True), ({"access": "no"}, False),
    ({"access": "customers"}, True)])
def test_roads_people_cannot_walk_along_are_left_out(tags, expected):
    assert osm.walkable(tags) is expected


def test_osm_tags_are_read_from_the_hstore_text():
    assert osm._tags('"shop"=>"supermarket","brand"=>"Aldi"') == {"shop": "supermarket", "brand": "Aldi"}
    assert osm._tags(None) == {}


# --- ABS mesh blocks -----------------------------------------------------------------------------
def counts_sheet():
    ws = Workbook().active
    ws.append(["Table 1. New South Wales (part 1)"])
    ws.append(["MB_CODE_2021", "MB_CATEGORY_NAME_2021", "AREA_ALBERS_SQKM", "Dwelling", "Person", "State"])
    ws.append(["10000010000", "Residential", 0.0209, 44, 63, 1])
    ws.append(["10000020000", "Education", 0.0148, 0, 0, 1])
    ws.append(["not a code", "Residential", 0.1, 5, 9, 1])
    return ws


def allocation_sheet():
    ws = Workbook().active
    ws.append(["MB_CODE_2021", "MB_CATEGORY_2021", "CHANGE_FLAG_2021", "AREA_ALBERS_SQKM", "SA1_CODE_2021",
               "SA2_CODE_2021", "SA2_NAME_2021", "GCCSA_CODE_2021"])
    ws.append(["10000010000", "Residential", 0, 0.02, "10102100701", "101021007", "Braidwood", "1RNSW"])
    ws.append(["11734650000", "Residential", 0, 0.02, "11703164401", "117031644", "Sydney (North) - Millers Point", "1GSYD"])
    return ws


def test_mesh_block_counts_are_read_by_heading():
    df = abs_sources.parse_mesh_counts_ws(counts_sheet())
    assert list(df["MB_CODE21"]) == [10000010000, 10000020000]
    assert list(df["dwellings"]) == [44, 0]
    assert list(df["persons"]) == [63, 0]
    assert df.loc[0, "category"] == "Residential"


def test_the_allocation_keeps_only_greater_sydney():
    df = abs_sources.parse_mesh_allocation_ws(allocation_sheet())
    assert list(df["MB_CODE21"]) == [11734650000]
    assert list(df["SA1_CODE21"]) == [11703164401]
    assert list(df["SA2_CODE21"]) == [117031644]


def test_a_sheet_without_the_expected_heading_is_refused():
    ws = Workbook().active
    ws.append(["some", "other", "sheet"])
    with pytest.raises(EtlError, match="MB_CODE_2021"):
        abs_sources.parse_mesh_counts_ws(ws)


# --- City of Sydney extras -----------------------------------------------------------------------
def test_city_of_sydney_columns_are_lower_cased_and_the_geometry_is_geom(tmp_path):
    from etl.core import Snapshot
    path = tmp_path / "stairs.geojson"
    gpd.GeoDataFrame({"OBJECTID": [1, 2], "No_Steps": [5, 9]}, geometry=[Point(151.2, -33.87), Point(151.21, -33.88)],
                     crs=4326).to_file(path, driver="GeoJSON")
    # Snapshot.path() is ROOT / path, and an absolute path replaces ROOT
    snap = Snapshot("cos_stairs", "x", "http://example.com", {"main": {"path": str(path)}}, retrieved_at="2026-09-25T00:00:00Z")
    out = cityofsydney.parse(snap)
    assert list(out) == ["cos_stairs"]
    assert list(out["cos_stairs"].columns) == ["objectid", "no_steps", "geom"]
    assert out["cos_stairs"].geometry.name == "geom"


def test_city_of_sydney_layers_with_too_few_rows_fail_the_check():
    gdf = gpd.GeoDataFrame({"objectid": [1, 2]}, geometry=[Point(151.2, -33.87), Point(151.21, -33.88)], crs=4326)
    issues = cityofsydney.check({"cos_stairs": gdf.rename_geometry("geom")}, None)
    assert any(level == "error" and "only 2 rows" in message for level, message in issues)
