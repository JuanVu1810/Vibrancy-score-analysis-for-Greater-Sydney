"""`python -m etl verify`: check the staged tables, and with --parity check the parsers against v1.

Parity: the v1 ABS files in data/ came from older ABS releases (June 2022 businesses, 2020-21 income,
the 2021 population release). Running today's parsers on those releases must give back the v1 numbers
exactly, which proves the parsers read the workbooks correctly.
"""
from __future__ import annotations

import zipfile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from . import quality
from .core import RAW, ROOT, Ctx
from .sources import SOURCES
from .sources import abs as abs_sources
from .stage import read_staged

ABS = abs_sources.ABS
PARITY = {
    "business": (f"{abs_sources.CABEE}/jul2018-jun2022/816509.xlsx", "cabee_jun2022_816509.xlsx"),
    "income": (f"{ABS}/statistics/labour/earnings-and-working-conditions/personal-income-australia/2020-21/"
               "Table%201%20-%20Total%20income%2C%20earners%20and%20summary%20statistics%20by%20geography%2C"
               "%202016-17%20to%202020-21.xlsx", "personal_income_2020-21_table1.xlsx"),
    "population": (f"{ABS}/statistics/people/population/regional-population-age-and-sex/2021/"
                   "32350DS0001_2021.xlsx", "erp_sa2_2021.xlsx"),
}


class Report:
    def __init__(self):
        self.failures = []

    def check(self, ok: bool, label: str, detail: str = "") -> None:
        print(f"  [{'PASS' if ok else 'FAIL'}] {label}" + (f"  ({detail})" if detail else ""))
        if not ok:
            self.failures.append(label)


def _cached(ctx: Ctx, url: str, filename: str):
    """Download an older release once into raw/parity/ (kept out of the main manifest)."""
    dest = RAW / "parity" / filename
    if not dest.exists() or not zipfile.is_zipfile(dest):  # missing, or left half-finished earlier
        ctx.log(f"  downloading {filename}")
        ctx.fetch_to(url, dest)
    return dest


def verify_staged(rep: Report) -> None:
    ctx = Ctx(verbose=False)
    print("Checking the staged tables (built-in checks and pandera schemas)")
    for src in SOURCES:
        out = {t: read_staged(t) for t in src.outputs}
        if any(v is None for v in out.values()):
            print(f"  [SKIP] {src.id}: not staged yet")
            continue
        for r in quality.assess(src, out, ctx):
            if r.status == "WARNING":
                print(f"  [WARN] {r.dataset}: {r.notes}")
            rep.check(r.status != "FAIL", f"{r.dataset} passes", r.notes if r.status == "FAIL" else "")


def _same(rep: Report, label: str, got: pd.DataFrame, want: pd.DataFrame, keys: list, tol: int = 0) -> None:
    got = got.sort_values(keys).reset_index(drop=True)
    want = want.sort_values(keys).reset_index(drop=True)
    rep.check(len(got) == len(want), f"{label}: same number of rows", f"parser {len(got)}, v1 {len(want)}")
    if len(got) != len(want) or list(got.columns) != list(want.columns):
        rep.check(list(got.columns) == list(want.columns), f"{label}: same columns")
        return
    rep.check(bool((got[keys].to_numpy() == want[keys].to_numpy()).all()), f"{label}: same keys")
    for col in [c for c in got.columns if c not in keys]:
        if pd.api.types.is_numeric_dtype(want[col]):
            a = got[col].astype("float64").to_numpy()
            b = want[col].astype("float64").to_numpy()
            diff = np.nanmax(np.abs(a - b)) if len(a) else 0
            same_nan = bool((np.isnan(a) == np.isnan(b)).all())
            rep.check(same_nan and diff <= tol, f"{label}: column {col}", f"largest difference {diff:g}")
        else:
            same = got[col].astype(str).str.strip() == want[col].astype(str).str.strip()
            rep.check(bool(same.all()), f"{label}: column {col}", f"{int((~same).sum())} differ")


def verify_parity(rep: Report) -> None:
    ctx = Ctx(verbose=True)
    print("\nParity: today's parsers on the older releases behind the v1 files")

    # Businesses: v1 is June 2022 (Jul 2018 to Jun 2022 release, sheet '2022')
    path = _cached(ctx, *PARITY["business"])
    ws = load_workbook(path, read_only=True, data_only=True)["2022"]
    got, _ = abs_sources.parse_business_ws(ws)
    want = pd.read_csv(ROOT / "data" / "Businesses.csv", encoding="utf-8-sig").rename(columns={"sa2_code": "SA2_CODE21"})
    _same(rep, "businesses (June 2022)", got, want[got.columns], ["SA2_CODE21", "industry_code"])

    # Income: v1 is 2020-21 total income (Table 1.4)
    path = _cached(ctx, *PARITY["income"])
    got, year = abs_sources.parse_income_ws(load_workbook(path, read_only=True, data_only=True)["Table 1.4"], "2020-21")
    want = pd.read_csv(ROOT / "data" / "Income.csv", encoding="utf-8-sig", na_values=["np"]).rename(columns={"sa2_code21": "SA2_CODE21"})
    _same(rep, "income (2020-21)", got, want[got.columns], ["SA2_CODE21"])

    # Population: v1 is the 2021 release (30 June 2021 estimates as first published), Table 3, persons
    path = _cached(ctx, *PARITY["population"])
    wb = load_workbook(path, read_only=True, data_only=True)
    got = abs_sources.parse_population_ws(abs_sources.population_sheet(wb))
    want = pd.read_csv(ROOT / "data" / "Population.csv", encoding="utf-8-sig").rename(columns={"sa2_code": "SA2_CODE21"})
    _same(rep, "population (2021 release)", got, want[got.columns], ["SA2_CODE21"])


def run(parity: bool) -> int:
    rep = Report()
    verify_staged(rep)
    if parity:
        verify_parity(rep)
    print()
    if rep.failures:
        print(f"RESULT: {len(rep.failures)} check(s) failed: " + "; ".join(rep.failures))
        return 1
    print("RESULT: all checks passed")
    return 0
